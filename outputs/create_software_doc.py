from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path('outputs/manuscript'); out.mkdir(parents=True,exist_ok=True)
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.6); sec.left_margin=sec.right_margin=Inches(.6)
normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(5)
for s in ['Title','Heading 1']:
 doc.styles[s].font.name='Arial'; doc.styles[s].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(15)
p=doc.add_paragraph('Software Specifications and Estimated Costs','Title'); p.alignment=1
p=doc.add_paragraph('Senior Citizen Assistance System for Barangay Services'); p.alignment=1
p=doc.add_paragraph('The table lists the software used for development and the proposed hosting and SMS services for deployment. Paid services are budget estimates.'); p.paragraph_format.space_after=Pt(8)
rows=[
['VS Code','Code editor','Free'],
['Git & GitHub','Version control and source code hosting','Open-Source / Free Plan'],
['XAMPP 8.2.12','Local development environment','Free / Open-Source'],
['Apache 2.4.58','Local web server','Open-Source'],
['PHP 8.2.12','Server-side programming language','Open-Source'],
['MariaDB 10.4.32','Relational database management system','Open-Source'],
['phpMyAdmin 5.2.1','Database administration tool','Open-Source'],
['HTML5','Web page structure and content','Open Web Standard'],
['CSS','Page styling and layout','Open Web Standard'],
['JavaScript','Interactive interface and asynchronous requests','Open Web Standard'],
['Bootstrap 5.3.3','Responsive front-end framework','Open-Source'],
['Bootstrap Icons 1.11.3','Interface icon library','Open-Source'],
['Chart.js 4.4.1','Dashboard charts','Open-Source'],
['Custom PHP Test Scripts','Application and HTTP endpoint testing','Locally Developed']]
rows=[r+['₱0.00','₱0.00'] for r in rows]
rows += [['SMTP Email Service','Verification and password-recovery emails','Email Account / SMTP','To be confirmed','To be confirmed'],['Textbee API — Scale','Planned SMS gateway','Monthly Subscription','₱1,874.76 / month','₱22,497.18'],['Hostinger.com — Cloud Startup','Planned application and database hosting','Subscription','₱1,749.74 / month','₱20,996.87'],['Estimated Paid Subtotal','Hosting and SMS subscriptions only','','₱3,624.50 / month','₱43,494.05']]
t=doc.add_table(rows=2, cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
widths=[1.4,2.1,1.4,1.2,1.2]
for c,w in zip(t.columns,widths): c.width=Inches(w)
t.cell(0,0).merge(t.cell(0,4)).text='Software Cost'
for c,txt in zip(t.rows[1].cells,['Item','Description','Access Mode','Unit Cost','Annual Cost']): c.text=txt
for vals in rows:
 for c,txt,w in zip(t.add_row().cells,vals,widths): c.text=txt; c.width=Inches(w)
for i,row in enumerate(t.rows):
 trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
 if i<2: trpr.append(OxmlElement('w:tblHeader'))
 for j,c in enumerate(row.cells):
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  pr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side); e.set(qn('w:w'),'65'); e.set(qn('w:type'),'dxa'); margins.append(e)
  pr.append(margins)
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i<2 else (WD_ALIGN_PARAGRAPH.RIGHT if j>=3 else WD_ALIGN_PARAGRAPH.LEFT)
   for r in p.runs: r.font.size=Pt(9); r.bold=(i<2 or i==len(t.rows)-1)
borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'000000'); borders.append(e)
t._tbl.tblPr.append(borders)
doc.add_page_break()
doc.add_paragraph('Cost Basis and Software Notes','Heading 1')
notes=[
'The local environment uses XAMPP 8.2.12, which includes Apache 2.4.58, PHP 8.2.12, MariaDB 10.4.32, and phpMyAdmin 5.2.1. These are local installation versions; the hosting environment may use different versions.',
'GitHub costs assume the Free plan. HTML, CSS, and JavaScript are web technologies with no separate licensing cost. SMTP email costs remain subject to the selected email service.',
'Textbee Scale and Hostinger Cloud Startup are proposed services. Their inclusion does not indicate that they have been purchased or integrated.',
'Currency conversion uses the budgeting rate of US$1 = ₱62.513 from the preceding cost estimate. Monthly costs are US$29.99 × ₱62.513 = ₱1,874.76 for Textbee and US$27.99 × ₱62.513 = ₱1,749.74 for Hostinger. Annual estimates multiply the unrounded monthly conversions by 12.',
'The Hostinger estimate uses the US$27.99 regular monthly price selected for the budget. The supplied screenshot separately lists US$7.99 per month under a 48-month promotional term and US$25.99 per month at renewal. Actual billing depends on the selected checkout offer.',
'The supplied Textbee Scale screenshot lists up to 25,000 messages per month, up to 15 active devices, and no daily message limit. It also offers annual billing at US$159.89, approximately ₱9,995.20 per year at the same exchange rate. The table uses monthly billing for 12 months instead.',
'The subtotal covers only hosting and the Textbee subscription. SIM load or mobile-plan charges, email expenses, domain charges, taxes, payment fees, and operating-system licensing are excluded. Actual peso charges depend on the exchange rate applied at payment.'
]
for s in notes: doc.add_paragraph(s)
doc.add_paragraph('Reference Links','Heading 1')
for s in ['Hostinger Philippine pricing: https://www.hostinger.com/ph/web-hosting','Textbee pricing: https://textbee.dev/pricing','Exchange-rate reference: https://twelvedata.com/markets/371101/forex/usd-php/historical-data']:
 doc.add_paragraph(s)
doc.save(out/'Software_Specifications_and_Costs.docx')
print(out/'Software_Specifications_and_Costs.docx')
