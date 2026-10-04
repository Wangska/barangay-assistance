from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
cases = [
{
'name':'Login and Access Control','actor':'Administrator, Barangay Staff, Senior Citizen',
'goal':'Allow registered users to access system functions according to their role and account status.',
'pre':['The user has an existing account.','The user can access the login page.'],
'post':['The user is authenticated and directed to the permitted system area.','Successful login and logout are recorded in the audit log.','Logout terminates the session.'],
'flow':[
('The user opens the login page.','The system displays the email and password fields.'),
('The user enters credentials and clicks Login.','The system checks required fields, the attempt limit, credentials, and account status.'),
('The user submits valid credentials for an Active account.','The system starts a session, clears failed attempts, records the login, and directs the user to the role-based area.'),
('The user selects a protected function.','The system checks the role and account status before granting access.'),
('The user selects Logout.','The system records the logout, destroys the session, and redirects to the home page.')],
'alt':[
('2A: Required fields are empty','The system displays: "Please fill in all fields."'),
('2B: Login attempt limit reached','After five failed attempts, further attempts are blocked until the current 15-minute window expires. The counter is session-based and keyed by email and IP address.'),
('3A: Invalid credentials or ineligible account status','The system displays: "Invalid email or password." and records a failed attempt.'),
('3B: Senior account is Pending or Rejected','Valid credentials allow restricted access to application tracking and related permitted pages.')],
'exc':[
('User is not logged in','The system redirects protected-page requests to Login.'),
('User has insufficient permission','The system displays: "Access denied."'),
('Authenticated account becomes inactive','The system blocks access to protected functions.')]
},
{
'name':'Online Senior Registration','actor':'Senior Citizen Applicant',
'goal':'Allow eligible senior citizens to submit an account application and personal profile for review.',
'pre':['The applicant is at least 60 years old.','The applicant can access the form and provide required information.','The applicant has an email address not already registered.'],
'post':['A Senior Citizen account is created with Pending status.','The profile, consent timestamp, and application submission are recorded.','Active barangay staff receive a registration notification.'],
'flow':[
('The applicant selects Create Account.','The system displays the senior registration form.'),
('The applicant enters location and personal information.','The form accepts address, name, birthdate, birthplace, sex, civil status, blood type, religion, and education.'),
('The applicant enters contact, benefits, and emergency details.','The form accepts contact numbers, benefits information, emergency contacts, and optional medical notes.'),
('The applicant selects an ID file, creates and confirms a password, and accepts consent.','The form includes the file, credentials, and consent in the submission.'),
('The applicant clicks Submit Registration.','The system validates required fields, age, email, mobile numbers, password, and upload.'),
('The applicant submits valid information.','The system saves the Pending account and profile, records the application and ID requirement, notifies active staff, and redirects to Login with confirmation.')],
'alt':[
('5A: Required information or consent is missing','The system requests completion of all required registration fields.'),
('5B: Applicant is below 60 or birthdate is invalid','The system rejects registration and explains the minimum age requirement.'),
('5C: Contact details are invalid','The system requests a valid email and Philippine mobile numbers.'),
('5D: Passwords do not match or are too short','The system requires matching passwords of at least eight characters.')],
'exc':[
('Email is already registered','Registration fails. The system asks the applicant to check their details or use a different email address.'),
('Identification upload fails','The system does not complete registration and displays a failure message.'),
('Database operation fails','The system rolls back the registration transaction and displays a failure message.')]
},
{
'name':'Record Senior Benefits Information','actor':'Senior Citizen Applicant',
'goal':'Record government membership identifiers, employment status, classification, and pension information as part of registration.',
'pre':['The applicant has opened the registration form.','The applicant is completing Contact and Benefits Information.'],
'post':['Benefits information is saved with the senior profile when registration succeeds.','Optional information may remain unspecified.'],
'flow':[
('The applicant opens Contact and Benefits Information.','The system displays membership, employment, classification, and pension fields.'),
('The applicant enters available GSIS, SSS, TIN, and PhilHealth numbers.','The form accepts these optional identifiers.'),
('The applicant selects employment status.','The form offers Employed, Self-Employed, Unemployed, and Retired.'),
('The applicant selects senior classification.','The form offers Indigent, Pensioner, and Supported.'),
('The applicant enters a monthly pension, if applicable.','The form accepts the amount in Philippine pesos.'),
('The applicant completes and submits registration.','The system checks that employment status and classification are supplied and saves the benefits information if registration succeeds.')],
'alt':[
('2A: Government identifiers are unavailable','The applicant may leave these optional fields blank.'),
('5A: Monthly pension is not provided','The system stores the pension amount as unspecified.'),
('6A: Employment status or classification is missing','The system rejects registration and requests completion of the required fields.')],
'exc':[
('Another registration section fails validation','Benefits information is not saved as a completed registration.'),
('Database operation fails','The system rolls back the registration transaction and displays a failure message.')]
},
{
'name':'Identification Upload','actor':'Senior Citizen Applicant',
'goal':'Allow the applicant to submit an identification document as a supporting registration requirement.',
'pre':['The applicant is completing registration.','The applicant has a JPG, PNG, or PDF file no larger than 2 MB.'],
'post':['The accepted file is stored and linked to the senior profile.','The identification requirement is recorded as Submitted when registration succeeds.'],
'flow':[
('The applicant locates Valid Identification.','The system displays the upload control, allowed formats, and size limit.'),
('The applicant selects an identification file.','The form attaches the selected file for submission.'),
('The applicant submits the registration form.','The system checks upload status, file size, and actual MIME type.'),
('The applicant supplies a valid file.','The system generates a random filename and saves the file in the configured identification storage directory.'),
('The applicant completes registration successfully.','The system links the file to the senior profile and records the requirement as Submitted.')],
'alt':[
('2A: No file is selected','The browser normally prevents submission. If browser validation is bypassed, the current server code permits registration and records the requirement as Missing.'),
('3A: File exceeds 2 MB','The system rejects the upload, and registration does not complete.'),
('3B: File format is unsupported','The system rejects the upload, and registration does not complete.')],
'exc':[
('File transfer fails','The system rejects the upload and displays a registration failure message.'),
('File cannot be saved to storage','The system stops registration and displays a failure message.'),
('Database operation fails after upload','Registration is rolled back. The file may remain in storage because the current code does not remove it during rollback.')]
}]
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11.7)
s.top_margin=s.bottom_margin=Inches(.55)
s.left_margin=s.right_margin=Inches(.7)
normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.line_spacing=1.05
style=doc.styles['Title']; style.font.name='Arial';style.font.size=Pt(11);style.font.bold=True;style.font.color.rgb=RGBColor(0,0,0)
style.paragraph_format.space_after=Pt(8)

def format_table(t,widths):
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
 pr.append(borders)
 for row in t.rows:
  tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  for c,w in zip(row.cells,widths):
   c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
   c._tc.get_or_add_tcPr().append(margins)

def fill(c,text,bullet=False):
 c.text=''
 for i,line in enumerate(text if isinstance(text,list) else [text]):
  p=c.paragraphs[0] if i==0 else c.add_paragraph()
  p.add_run(('\u2022  ' if bullet else '')+line)
  p.paragraph_format.space_after=Pt(3 if bullet else 0)

def headingrow(t,text):
 row=t.add_row();c=row.cells[0].merge(row.cells[-1]);fill(c,text)
 p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].bold=True
 p.paragraph_format.keep_with_next=True
 return row

for n,case in enumerate(cases,1):
 if n>1:doc.add_page_break()
 doc.add_paragraph('Use-Case Narrative',style='Title')
 t=doc.add_table(rows=6,cols=2)
 values=[('Use Case ID',f'UC-SCAS-{n:03}'),('Use Case Name',case['name']),('Actor',case['actor']),('Goal',case['goal']),('Preconditions',case['pre']),('Postconditions',case['post'])]
 for row,(label,value) in zip(t.rows,values):
  fill(row.cells[0],label);fill(row.cells[1],value,isinstance(value,list))
 format_table(t,[1.55,5.55])
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(8);p.add_run(' ').font.size=Pt(3)
 t=doc.add_table(rows=0,cols=2)
 headingrow(t,'MAIN FLOW')
 row=t.add_row()
 for c,txt in zip(row.cells,['Actor Action','System Response']):
  fill(c,txt);c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER;c.paragraphs[0].runs[0].bold=True
  c.paragraphs[0].paragraph_format.keep_with_next=True
 for i,(a,b) in enumerate(case['flow'],1):
  row=t.add_row();fill(row.cells[0],f'{i}. {a}');fill(row.cells[1],b)
 for key,title in [('alt','ALTERNATIVE FLOW OF EVENTS'),('exc','EXCEPTIONS')]:
  headingrow(t,title)
  for label,response in case[key]:
   row=t.add_row();c=row.cells[0].merge(row.cells[1]);c.text=''
   p=c.paragraphs[0];r=p.add_run(label);r.italic=True
   p.add_run('\n- '+response)
 format_table(t,[3.25,3.85])
path=Path('outputs/use-case-narratives/Senior_Citizen_Use_Case_Narratives.docx')
doc.save(path)
print(path.resolve())
