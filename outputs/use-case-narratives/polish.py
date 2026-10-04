from docx import Document
from docx.oxml.ns import qn
p='outputs/use-case-narratives/Senior_Citizen_Use_Case_Narratives.docx'
d=Document(p)
for root in [d._element,d.styles.element]:
 for node in list(root.iter(qn('w:pBdr'))): node.getparent().remove(node)
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   if c.text in ['MAIN FLOW','Actor Action','System Response','ALTERNATIVE FLOW OF EVENTS','EXCEPTIONS']:
    for paragraph in c.paragraphs:
     for r in paragraph.runs: r.bold=True; r.font.cs_bold=True
     
d.save(p)
