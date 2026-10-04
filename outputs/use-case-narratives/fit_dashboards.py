from docx import Document
from docx.oxml.ns import qn
p='outputs/use-case-narratives/Dashboard_Use_Case_Narratives.docx'
d=Document(p)
for t in d.tables[-2:]:
 for node in t._tbl.iter(qn('w:tcMar')):
  for edge in ['top','bottom']:
   e=node.find(qn('w:'+edge))
   if e is not None:e.set(qn('w:w'),'30')
 for row in t.rows:
  for c in row.cells:
   for para in c.paragraphs:para.paragraph_format.line_spacing=1

d.save(p)
